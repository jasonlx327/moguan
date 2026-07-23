import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT.parent / "docs" / "墨观产品定义框架与愿景.md"
OUT = ROOT.parent / "docs" / "墨观产品定义框架与愿景.docx"

# standard_business_brief preset with a named CJK font override.
FONT = "PingFang SC"
BODY_SIZE = 11
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
BODY = "1F2937"
MUTED = "5B6573"
LIGHT = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM_DXA = 80
CELL_START_END_DXA = 120


def set_run_font(run, *, size=BODY_SIZE, color=BODY, bold=None, italic=None):
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        r_fonts.set(qn(slot), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline_runs(paragraph, text, *, size=BODY_SIZE, color=BODY, base_bold=False):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_run_font(paragraph.add_run(part[2:-2]), size=size, color=color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            set_run_font(paragraph.add_run(part[1:-1]), size=max(9, size - 1), color=INK)
        else:
            set_run_font(paragraph.add_run(part), size=size, color=color, bold=base_bold)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def paragraph_left_border(paragraph, color=BLUE, size="18", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True


def set_header_footer(section):
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("墨观｜产品定义、总体框架与长期愿景"), size=9, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    set_run_font(p.add_run("第 "), size=9, color=MUTED)
    add_field(p, "PAGE")
    set_run_font(p.add_run(" 页"), size=9, color=MUTED)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("产品蓝图"), size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("墨观"), size=34, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("产品定义、总体框架与长期愿景"), size=17, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("基于产品设定 v0.5｜2026 年 7 月 23 日"), size=10, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.52)
    p.paragraph_format.right_indent = Inches(0.52)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(30)
    p.paragraph_format.line_spacing = 1.25
    shade_paragraph(p, CALLOUT)
    paragraph_left_border(p)
    add_inline_runs(
        p,
        "从中国古代星空出发，结合真实天象、地球天气与公共事件，"
        "用先秦至宋代的分层文献体系分析世界如何变化。",
        size=13,
        color=INK,
        base_bold=True,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(54)
    set_run_font(p.add_run("仰观天象｜俯察地理｜观察人事｜验证判断"), size=10, color=MUTED)

    doc.add_page_break()


def add_numbering_definition(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    abstract_id = max(abstract_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.extend([start, num_fmt, lvl_text, suff])

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, ind])
    lvl.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        r_fonts.set(qn(slot), FONT)
    r_pr.append(r_fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    return abstract_id


def new_num_instance(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    values = {
        "top": CELL_TOP_BOTTOM_DXA,
        "bottom": CELL_TOP_BOTTOM_DXA,
        "start": CELL_START_END_DXA,
        "end": CELL_START_END_DXA,
    }
    for side, value in values.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, rows):
    col_count = len(rows[0])
    if col_count == 3:
        widths = [2050, 3500, 3810]
    elif col_count == 2:
        widths = [2450, 6910]
    else:
        widths = [TABLE_WIDTH_DXA // col_count] * col_count
        widths[-1] += TABLE_WIDTH_DXA - sum(widths)

    table = doc.add_table(rows=len(rows), cols=col_count)
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])

    for r_idx, source_row in enumerate(rows):
        for c_idx, text in enumerate(source_row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            if r_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shade_cell(cell, LIGHT)
                add_inline_runs(p, text, size=10, color=INK, base_bold=True)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                add_inline_runs(p, text, size=10, color=BODY, base_bold=(c_idx == 0))

    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)
    return table


def strip_markdown_cell(text):
    return text.strip().strip("|").strip()


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [strip_markdown_cell(cell) for cell in lines[index].strip().split("|")[1:-1]]
        rows.append(cells)
        index += 1
    if len(rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        return None, start
    return [rows[0]] + rows[2:], index


def add_body(doc, markdown):
    body = markdown.split("---", 1)[1] if "---" in markdown else markdown
    lines = body.splitlines()
    bullet_abstract = add_numbering_definition(doc, "bullet")
    decimal_abstract = add_numbering_definition(doc, "decimal")
    index = 0

    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("|"):
            table_rows, next_index = parse_table(lines, index)
            if table_rows:
                add_table(doc, table_rows)
                index = next_index
                continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline_runs(p, heading.group(2), size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, base_bold=True)
            index += 1
            continue

        if stripped.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.right_indent = Inches(0.22)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(9)
            p.paragraph_format.line_spacing = 1.20
            shade_paragraph(p, CALLOUT)
            paragraph_left_border(p)
            add_inline_runs(p, " ".join(quote_lines), size=11.5, color=INK, base_bold=True)
            continue

        bullet = re.match(r"^-\s+(.*)$", stripped)
        if bullet:
            num_id = new_num_instance(doc, bullet_abstract)
            while index < len(lines):
                match = re.match(r"^-\s+(.*)$", lines[index].strip())
                if not match:
                    break
                p = doc.add_paragraph()
                apply_num(p, num_id)
                add_inline_runs(p, match.group(1))
                index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            num_id = new_num_instance(doc, decimal_abstract)
            while index < len(lines):
                match = re.match(r"^\d+\.\s+(.*)$", lines[index].strip())
                if not match:
                    break
                p = doc.add_paragraph()
                apply_num(p, num_id)
                add_inline_runs(p, match.group(1))
                index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", ">", "|", "- "))
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.10
        add_inline_runs(p, " ".join(paragraph_lines))


def set_core_properties(doc):
    props = doc.core_properties
    props.title = "墨观｜产品定义、总体框架与长期愿景"
    props.subject = "墨观产品蓝图"
    props.author = "墨观"
    props.keywords = "墨观, 中国古典世界分析, 产品定义, 愿景, 产品框架"
    props.comments = "基于墨观产品设定 v0.5"


def main():
    doc = Document()
    section = doc.sections[0]
    configure_page(section)
    configure_styles(doc)
    set_header_footer(section)
    set_core_properties(doc)
    add_cover(doc)
    add_body(doc, SOURCE.read_text(encoding="utf-8"))
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
